import os
import io
import docx
from pypdf import PdfWriter, PdfReader
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors

def create_markdown_file(path, title, policy_id, version, status, jurisdiction, content):
    md_content = f"""# {title}

**Policy ID**: {policy_id}
**Version**: {version}
**Status**: {status}
**Jurisdiction**: {jurisdiction}
**Effective Date**: 2026-01-01
**Owner**: Human Resources

{content}
"""
    with open(path, 'w', encoding='utf-8') as f:
        f.write(md_content)

def create_docx_file(path, title, policy_id, version, content):
    doc = docx.Document()
    doc.add_heading(title, 0)
    
    doc.add_paragraph(f"Policy ID: {policy_id}")
    doc.add_paragraph(f"Version: {version}")
    doc.add_paragraph("Status: active")
    doc.add_paragraph("Jurisdiction: Global")
    doc.add_paragraph("Effective Date: 2026-01-01")
    doc.add_paragraph("Owner: Finance")
    
    for section_title, section_text in content:
        doc.add_heading(section_title, level=1)
        doc.add_paragraph(section_text)
        
    doc.save(path)

def create_pdf_file(path, title, policy_id, version, content):
    packet = io.BytesIO()
    can = canvas.Canvas(packet, pagesize=letter)
    can.setFont("Helvetica-Bold", 16)
    can.drawString(50, 750, title)
    
    can.setFont("Helvetica", 12)
    can.drawString(50, 720, f"Policy ID: {policy_id}")
    can.drawString(50, 700, f"Version: {version}")
    can.drawString(50, 680, "Status: active")
    can.drawString(50, 660, "Jurisdiction: US")
    can.drawString(50, 640, "Effective Date: 2026-01-01")
    can.drawString(50, 620, "Owner: Finance")
    
    y = 580
    for section_title, section_text in content:
        can.setFont("Helvetica-Bold", 14)
        if y < 100:
            can.showPage()
            y = 750
        can.drawString(50, y, section_title)
        y -= 20
        
        can.setFont("Helvetica", 11)
        # Simple text wrapping for PDF
        words = section_text.split()
        line = ""
        for word in words:
            if can.stringWidth(line + word, "Helvetica", 11) < 500:
                line += word + " "
            else:
                if y < 50:
                    can.showPage()
                    y = 750
                    can.setFont("Helvetica", 11)
                can.drawString(50, y, line)
                y -= 15
                line = word + " "
        if line:
            if y < 50:
                can.showPage()
                y = 750
                can.setFont("Helvetica", 11)
            can.drawString(50, y, line)
            y -= 25
            
    can.save()
    
    packet.seek(0)
    with open(path, "wb") as f:
        f.write(packet.read())

def main():
    os.makedirs("data/policies", exist_ok=True)
    
    # MD 1: WFH (Active)
    wfh_content = """## 1. Purpose
This policy establishes the guidelines and expectations for remote work, telecommuting, and work-from-home (WFH) arrangements to ensure operational efficiency, data security, and employee well-being.

## 2. Scope and Eligibility
This policy applies to all full-time employees globally, with specific departmental exceptions.
- **Engineering and Product**: Eligible for fully remote work or up to 4 days WFH per week.
- **Sales and Marketing**: Eligible for up to 3 days WFH per week. Client-facing days do not count against this quota.
- **Customer Support**: Must adhere to a strict hybrid schedule (Monday, Wednesday, Friday in-office) due to compliance and secure network requirements. Fully remote work is not permitted for Support unless a formal exception is granted by the VP of Support.
- **Probationary Employees**: New hires in their 90-day probationary period are only eligible for 1 day of WFH per week across all departments.

## 3. Remote Work Allowances and Hardware
Employees engaging in WFH are eligible for a one-time Home Office Stipend of $500 to purchase ergonomic equipment (chairs, desks). 
- Hardware (laptops, monitors) will be directly provided and provisioned by the IT Department. Do not purchase hardware using the stipend.
- For internet and mobile data reimbursement, refer to the Expense Reimbursement Policy (POL-EXPENSE-04). The company provides a flat $50 monthly allowance for home internet, which is automatically added to the paycheck.

## 4. Work Environment and Security
Remote employees must ensure a private, distraction-free environment. 
All company work must be conducted on secure, company-issued devices using the corporate VPN. The use of public Wi-Fi without the corporate VPN is strictly prohibited and constitutes a security violation.

## 5. Sick Leave During Remote Work
If an employee falls ill on a scheduled WFH day, they must formally request sick leave rather than attempting to work remotely at diminished capacity. See POL-LEAVE-02 for sick leave accrual and usage guidelines.
"""
    create_markdown_file(
        "data/policies/POL-WFH-01.md",
        "Work From Home Policy",
        "POL-WFH-01",
        "2.0",
        "active",
        "Global",
        wfh_content
    )
    
    # MD 2: Sick Leave (Active)
    leave_content = """## 1. Purpose
This document details the Paid Time Off (PTO), Sick Leave, and specialized leave categories available to employees to promote work-life balance and health.

## 2. Sick Leave Accrual and Usage
Full-time employees accrue 1.25 sick days per month, totaling 15 days annually. 
- Sick leave can be carried over to the subsequent calendar year up to a maximum cap of 30 days. Any accrued days beyond 30 will be forfeited at the end of the calendar year.
- If an absence due to illness exceeds 3 consecutive business days, the employee must provide a medical certificate or doctor's note to Human Resources upon their return.
- Mental health days are classified under sick leave and do not require prior managerial approval if taken as single days.

## 3. Paid Time Off (PTO)
Employees receive 20 days of PTO per year, front-loaded on January 1st. 
- PTO requests must be submitted at least two weeks in advance for vacations exceeding 5 days.
- Unlike sick leave, unused PTO is paid out upon termination of employment.

## 4. Bereavement and Jury Duty
- **Bereavement**: Employees are granted up to 5 days of paid bereavement leave in the event of the loss of an immediate family member.
- **Jury Duty**: The company provides up to 10 days of paid leave for mandatory jury duty. Employees must submit their jury summons to HR.

## 5. Overlap with WFH
If an employee is scheduled to work from home (WFH) but is unable to work due to illness, they must log the day as Sick Leave. Working a partial day while sick is generally discouraged unless formally coordinated as a half-day with the manager.
"""
    create_markdown_file(
        "data/policies/POL-LEAVE-02.md",
        "Paid Time Off and Leave Policy",
        "POL-LEAVE-02",
        "2.0",
        "active",
        "Global",
        leave_content
    )
    
    # MD 3: Sick Leave (Superseded - deliberately conflicts with Active version)
    leave_v1_content = """## 1. Purpose
This document details the leave available to employees.

## 2. Sick Leave
Employees accrue 0.5 sick days per month, totaling 6 days annually. 
- Sick leave **cannot** be carried over to the next year. It is a "use it or lose it" system.
- No doctor's note is ever required for sick leave, regardless of duration.

## 3. Paid Time Off (PTO)
Employees accrue 1 day of PTO per month. PTO does not pay out upon termination.
"""
    create_markdown_file(
        "data/policies/POL-LEAVE-02-v1.md",
        "Paid Time Off and Leave Policy",
        "POL-LEAVE-02",
        "1.0",
        "superseded",
        "Global",
        leave_v1_content
    )
    
    # DOCX: Travel (Active)
    travel_content = [
        ("1. Purpose", "The Corporate Travel Policy governs all business-related travel to ensure safety, cost-effectiveness, and compliance with corporate finance standards."),
        ("2. Booking Flights and Accommodation", "All travel must be booked through the corporate travel portal (Concur). Direct bookings on airline or hotel websites will not be reimbursed unless an emergency exception is granted.\n- Standard Employees must book Economy class for all flights under 8 hours. Flights over 8 hours are eligible for Premium Economy.\n- Executives (VP level and above) are permitted to book Business Class for any flight exceeding 4 hours."),
        ("3. Daily Allowance (Per Diem)", "Employees traveling for business are granted a daily allowance (per diem) to cover meals and incidentals.\n- Domestic Travel (US): $75 per day.\n- International Travel: $120 per day.\nNo itemized receipts are required for per diem claims. However, if a meal is paid for by a client or another executive, the employee must deduct $25 from their daily per diem claim."),
        ("4. Client-Billable Travel", "If travel expenses are being billed directly to a client, employees must adhere to the client's specific travel policy, which supersedes this corporate policy. For instance, if the client requires itemized receipts for meals, the per diem model is suspended and the Expense Reimbursement Policy (POL-EXPENSE-04) applies.")
    ]
    create_docx_file("data/policies/POL-TRAVEL-03.docx", "Corporate Travel Policy", "POL-TRAVEL-03", "2.1", travel_content)
    
    # PDF: Expense (Active)
    expense_content = [
        ("1. General Guidelines", "This policy outlines the procedure for reimbursing employees for out-of-pocket expenses incurred during the course of business operations."),
        ("2. Submission Deadlines", "All expense reports must be submitted within 30 days of the transaction date. Expenses submitted after 60 days will be permanently denied without exception. Reports must be filed via the Expensify portal."),
        ("3. Receipt Requirements", "Itemized receipts are strictly mandatory for any individual expense exceeding $25. Credit card summary slips that do not show the itemized breakdown are unacceptable and will be rejected by Finance."),
        ("4. Overlap with Other Policies", "- Travel Meals: Standard daily meals during travel are covered by the Travel Policy (POL-TRAVEL-03) per diem and do not require receipts. However, business meals involving client entertainment must be submitted as expenses here, requiring itemized receipts and a list of attendees.\n- Hardware: Keyboards, mice, and home-office equipment must be expensed using the WFH Home Office Stipend (POL-WFH-01). Laptops and monitors cannot be expensed; they must be procured through IT."),
        ("5. Prohibited Expenses", "The company will not reimburse: alcohol (unless part of an approved client dinner), parking tickets, traffic violations, personal grooming, or in-flight entertainment.")
    ]
    create_pdf_file("data/policies/POL-EXPENSE-04.pdf", "Expense Reimbursement Policy", "POL-EXPENSE-04", "1.5", expense_content)
    
    print("Realistic expanded corpus generated successfully.")

if __name__ == "__main__":
    main()
