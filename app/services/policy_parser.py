import os
import re
from typing import List, Dict, Any, Optional
from app.schemas.rag import ParsedDocument, PolicyMetadata, ParsedSection

class PolicyParser:
    """
    Format-aware parser for Markdown, DOCX, and PDF policies.
    """
    
    def parse_file(self, file_path: str) -> ParsedDocument:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".md":
            return self._parse_markdown(file_path)
        elif ext == ".docx":
            return self._parse_docx(file_path)
        elif ext == ".pdf":
            return self._parse_pdf(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    def _parse_markdown(self, file_path: str) -> ParsedDocument:
        with open(file_path, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        return self._extract_from_lines(file_path, lines)

    def _parse_docx(self, file_path: str) -> ParsedDocument:
        import docx
        doc = docx.Document(file_path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        return self._extract_from_lines(file_path, lines)

    def _parse_pdf(self, file_path: str) -> ParsedDocument:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        lines = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                lines.extend(text.split('\n'))
        return self._extract_from_lines(file_path, lines)

    def _extract_from_lines(self, file_path: str, lines: List[str]) -> ParsedDocument:
        """
        Custom parsing logic that handles both Markdown and normalized plain text (from PDF/DOCX)
        assuming a standard convention for this corpus.
        """
        title = "Unknown Policy"
        meta_dict = {
            "policy_id": "UNKNOWN",
            "version": "1.0",
            "status": "active",
            "jurisdiction": "Global",
            "effective_from": None,
            "owner": "Unknown"
        }
        
        sections: List[ParsedSection] = []
        current_section = None
        cleaned_lines = []
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Clean noise (simple example)
            if "page" in line.lower() and len(line) < 10:
                continue
                
            # Extract Markdown H1 or first line as title
            if line.startswith("# "):
                title = line[2:].strip()
                continue
            elif not title or title == "Unknown Policy":
                if not ":" in line and len(line) > 5 and len(line) < 100:
                    title = line
                    continue
                    
            # Extract Metadata
            meta_match = re.match(r"\**([A-Za-z\s]+)\**:\s*(.+)", line)
            if meta_match:
                key = meta_match.group(1).strip().lower().replace(" ", "_")
                val = meta_match.group(2).strip()
                if "date" in key:
                    key = "effective_from"
                if key in meta_dict:
                    meta_dict[key] = val
                    continue

            # Extract Section Headings (H2 or number format)
            # Markdown: ## 1. Purpose
            # Plain text: 1. Purpose
            section_match = re.match(r"^(?:##\s*)?(\d+\.\s+[A-Za-z\s]+)$", line)
            if section_match:
                heading = section_match.group(1)
                current_section = ParsedSection(heading=heading, level=2, content="")
                sections.append(current_section)
            else:
                if current_section:
                    current_section.content += line + "\n"
                cleaned_lines.append(line)

        # Build Metadata
        metadata = PolicyMetadata(
            policy_id=meta_dict.get("policy_id", "UNK"),
            title=title,
            policy_type="Standard",
            version=meta_dict.get("version", "1.0"),
            status=meta_dict.get("status", "active"),
            jurisdiction=meta_dict.get("jurisdiction"),
            effective_from=meta_dict.get("effective_from"),
            source=os.path.basename(file_path)
        )
        
        # Clean section contents
        for s in sections:
            s.content = s.content.strip()

        return ParsedDocument(
            source_id=os.path.basename(file_path),
            metadata=metadata,
            cleaned_content="\n".join(cleaned_lines),
            sections=sections
        )
